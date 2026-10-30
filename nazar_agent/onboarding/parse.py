"""A list of people as the owner has it → rows: name, site, how it was written.

The owner sends what they already have: a photo of the foreman's notebook,
an Excel sheet, a CSV or text pasted from a chat. A photo goes to Muse Spark
as an image; a sheet is turned into text first. The model only reads — it
reports what is written and marks what it could not read instead of
guessing. Matching against NAZAR (sites, existing employees) is done by code
afterwards (``review``).
"""

from __future__ import annotations

import base64
import csv
import io
from typing import Any

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from nazar_agent.engine import tool_schema
from nazar_agent.llm import LLMClient
from nazar_agent.usage import count

MAX_SHEET_ROWS = 500
MAX_TEXT = 20_000
IMAGE_TYPES = {"image/jpeg", "image/png", "image/webp"}
SHEET_TYPES = {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}

INSTRUCTIONS = """\
You read a list of employees that a business owner in Uzbekistan wants to add
to NAZAR, a time-and-attendance service: a photo of a handwritten notebook, a
spreadsheet or text from a chat. Names may be in Cyrillic or Latin, Russian
or Uzbek.

For every person on the list return one row:
- full_name: as written, in the same script; fix only capitalisation;
- site: the site (object, shop, building) the person belongs to when the
  list says so — e.g. under a heading "Объект 1 — Чиланзар" or in a column;
  empty when it does not say;
- as_written: the line exactly as you see it;
- legible: false when you are not sure you read the name right;
- note: why it is unclear, otherwise empty.

Do not invent people, names or sites. Ignore phone numbers, positions,
amounts and other columns. Skip headings, totals and empty lines.
"""


class ParseError(Exception):
    """The input cannot be read; ``said(lang)`` is the message for the person, in the window's language."""

    def __init__(self, kind: str) -> None:
        super().__init__(MESSAGES[kind]["ru"])
        self.kind = kind

    def said(self, lang: str) -> str:
        return MESSAGES[self.kind].get(lang, MESSAGES[self.kind]["en"])


MESSAGES = {
    "empty": {
        "ru": "Пришлите фото списка, Excel или вставьте текст.",
        "uz": "Ro'yxat suratini, Excel faylini yuboring yoki matnni qo'ying.",
        "en": "Send a photo of the list or an Excel file, or paste the text.",
    },
    "unreadable": {
        "ru": "Не получилось разобрать список — попробуйте ещё раз или пришлите текстом.",
        "uz": "Ro'yxatni o'qib bo'lmadi — qayta urinib ko'ring yoki matn bilan yuboring.",
        "en": "Could not read the list — try again or paste it as text.",
    },
    "format": {
        "ru": "Этот формат не подходит: пришлите фото (JPG, PNG), Excel (.xlsx) или CSV.",
        "uz": "Bu format mos emas: surat (JPG, PNG), Excel (.xlsx) yoki CSV yuboring.",
        "en": "This file does not fit: send a photo (JPG, PNG), an Excel file (.xlsx) or a CSV.",
    },
    "excel": {
        "ru": "Не получилось открыть Excel-файл.",
        "uz": "Excel faylini ochib bo'lmadi.",
        "en": "Could not open the Excel file.",
    },
    "encoding": {
        "ru": "Не получилось прочитать файл — сохраните его в UTF-8.",
        "uz": "Faylni o'qib bo'lmadi — uni UTF-8 da saqlang.",
        "en": "Could not read the file — save it as UTF-8.",
    },
}


class ListedPerson(BaseModel):
    model_config = ConfigDict(extra="forbid")

    full_name: str
    site: str
    as_written: str
    legible: bool
    note: str


class ListedPeople(BaseModel):
    model_config = ConfigDict(extra="forbid")

    people: list[ListedPerson] = Field(description="In the order of the list.")


PEOPLE_FORMAT: dict[str, Any] = {
    "type": "json_schema",
    "name": "ListedPeople",
    "schema": tool_schema(ListedPeople),
    "strict": True,
}


async def read_list(
    llm: LLMClient,
    model: str,
    *,
    text: str | None = None,
    file: bytes | None = None,
    content_type: str | None = None,
) -> list[ListedPerson]:
    content: list[dict[str, Any]] = [{"type": "input_text", "text": "Here is the list."}]
    if file is not None:
        content.append(_file_content(file, content_type or ""))
    if text and text.strip():
        content.append({"type": "input_text", "text": text.strip()[:MAX_TEXT]})
    if len(content) == 1:
        raise ParseError("empty")

    turn = await llm.respond(
        model=model,
        instructions=INSTRUCTIONS,
        input=[{"type": "message", "role": "user", "content": content}],
        tools=[],
        text_format=PEOPLE_FORMAT,
        role="import",
    )
    count("import", model, turn.input_tokens, turn.output_tokens)
    try:
        return ListedPeople.model_validate_json(turn.text).people
    except ValidationError as exc:
        raise ParseError("unreadable") from exc


def _file_content(file: bytes, content_type: str) -> dict[str, Any]:
    content_type = file_type(file, content_type)
    if content_type in IMAGE_TYPES:
        data = base64.b64encode(file).decode()
        return {"type": "input_image", "image_url": f"data:{content_type};base64,{data}"}
    if content_type in SHEET_TYPES:
        return {"type": "input_text", "text": sheet_text(file)}
    if content_type.startswith("text/"):
        return {"type": "input_text", "text": _decode(file)[:MAX_TEXT]}
    raise ParseError("format")


def file_type(file: bytes, declared: str) -> str:
    """What the file is by its first bytes, not by the name a browser or a server gave it:
    a server without the .xlsx type sends an Excel file as application/octet-stream."""
    if file.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if file.startswith(b"\x89PNG"):
        return "image/png"
    if file[:4] == b"RIFF" and file[8:12] == b"WEBP":
        return "image/webp"
    if file.startswith(b"PK\x03\x04"):  # a zip: .xlsx is one
        return next(iter(SHEET_TYPES))
    if declared in IMAGE_TYPES or declared in SHEET_TYPES or declared.startswith("text/"):
        return declared
    unnamed = declared in ("", "application/octet-stream")  # a CSV from a server that does not know .csv
    return "text/plain" if unnamed and file and b"\x00" not in file[:4096] else declared


def sheet_text(file: bytes) -> str:
    """Every non-empty row of every sheet, cells separated by " | "."""
    try:
        workbook = load_workbook(io.BytesIO(file), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many kinds for a broken file
        raise ParseError("excel") from exc
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if cells:
                lines.append(" | ".join(cells))
            if len(lines) >= MAX_SHEET_ROWS:
                break
    workbook.close()
    return "\n".join(lines)[:MAX_TEXT]


def _decode(file: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = file.decode(encoding)
        except UnicodeDecodeError:
            continue
        # A CSV reads better as "a | b | c" than with its delimiters.
        if any(delimiter in text.split("\n", 1)[0] for delimiter in ",;\t"):
            try:
                dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t")
            except csv.Error:
                return text
            return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text), dialect))
        return text
    raise ParseError("encoding")
